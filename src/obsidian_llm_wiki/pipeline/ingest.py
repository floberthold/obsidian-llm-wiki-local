"""
Ingest pipeline: raw note → chunk → analyze → embed → update state.

Uses the configured fast model (default: qwen3:4b) for analysis.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import logging
import os
import re
import time
from datetime import datetime
from pathlib import Path

from ..config import Config
from ..models import AnalysisResult, Concept, RawNoteRecord
from ..protocols import LLMClientProtocol
from ..state import StateDB
from ..structured_output import request_structured
from ..telemetry import emit_event
from ..vault import (
    chunk_text,
    generate_aliases,
    parse_note,
    sanitize_filename,
    sanitize_wikilink_target,
    write_note,
)

log = logging.getLogger(__name__)

_MAX_SOURCE_PATH_LEN = 220

_QUALITY_RANK: dict[str, int] = {"high": 2, "medium": 1, "low": 0}

_SYSTEM = (
    "You are a knowledge analyst. Read the provided note and extract structured information. "
    "Be concise and accurate. Do not invent information not present in the note. "
    "Detect the primary language of the note and return its ISO 639-1 code in the 'language' field "
    "(e.g. 'en', 'fr', 'de'). Use null if uncertain."
)


def _content_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _condensed_source_dir_name(relative_dir: Path) -> str:
    label = sanitize_filename(relative_dir.name, max_len=40) or "source"
    digest = hashlib.sha1(relative_dir.as_posix().encode("utf-8")).hexdigest()[:8]
    return f"{label}-{digest}"


def _source_summary_dir(config: Config, relative_dir: Path) -> Path:
    exact_dir = config.sources_dir / relative_dir
    if len(str(exact_dir)) <= _MAX_SOURCE_PATH_LEN:
        return exact_dir
    return config.sources_dir / _condensed_source_dir_name(relative_dir)


def _source_summary_path(config: Config, path: Path) -> Path:
    try:
        rel_path = path.relative_to(config.raw_dir)
    except ValueError:
        rel_path = path.relative_to(config.conversions_dir)
    exact_path = config.sources_dir / rel_path
    if len(str(exact_path)) <= _MAX_SOURCE_PATH_LEN:
        return exact_path

    if rel_path.parent == Path("."):
        stem = sanitize_filename(rel_path.stem, max_len=60)
        digest = hashlib.sha1(rel_path.as_posix().encode("utf-8")).hexdigest()[:8]
        return config.sources_dir / f"{stem}-{digest}{rel_path.suffix}"

    return _source_summary_dir(config, rel_path.parent) / rel_path.name


def _read_text_with_fallback(path: Path) -> str:
    """Read text files with practical encoding fallbacks for user-authored notes."""
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    # Latin-1 should always succeed, but keep a final safe guard.
    return raw.decode("utf-8", errors="replace")


def _build_analysis_prompt(
    body: str,
    existing_concepts: list[str],
    path_name: str = "",
    chunk_label: str = "",
) -> str:
    concepts_hint = ", ".join(existing_concepts[:30]) if existing_concepts else "none yet"
    label = f" {chunk_label}" if chunk_label else ""
    return (
        f"Analyze this note{label} and extract structured metadata.\n\n"
        f"Existing wiki concepts (reuse these names where applicable): {concepts_hint}\n\n"
        f"For each concept, provide 3-5 short surface forms used in running text "
        f"(abbreviations, short names). Example: name='Program Counter (PC)', "
        f"aliases=['PC', 'program counter']. Use empty list if no natural aliases exist.\n\n"
        f"DO NOT include the following as concepts — they are auto-generated context markers, not knowledge concepts:\n"
        f"- Standalone page labels (already filtered elsewhere)\n"
        f"- Bare dates, e.g. 'January 2024', '2024-03-15', '3/15/2024'\n"
        f"- Meeting or transcript timestamps, e.g. '00:05:23', '10:30 AM'\n"
        f"- Generic navigation entries such as 'index' or 'table of contents'\n\n"
        f"NOTE CONTENT:\n{body}"
    )


def _merge_chunk_results(results: list[AnalysisResult]) -> AnalysisResult:
    """Merge AnalysisResults from multiple chunks into one.

    Concepts and topics: union (deduplicated, insertion order preserved).
    Aliases for the same concept are merged across chunks.
    Summary: first chunk's (intro is most representative).
    Quality: minimum across chunks (conservative).
    """
    if len(results) == 1:
        return results[0]

    # Dedup concepts by canonical name (case-insensitive), merge aliases
    seen: dict[str, list[str]] = {}  # lower(name) -> accumulated aliases
    order: list[str] = []  # canonical names in insertion order
    canonical_by_lower: dict[str, str] = {}

    for r in results:
        for c in r.concepts:
            key = c.name.lower()
            if key not in seen:
                seen[key] = list(c.aliases)
                order.append(key)
                canonical_by_lower[key] = c.name
            else:
                existing_lower = {a.lower() for a in seen[key]}
                for a in c.aliases:
                    if a.lower() not in existing_lower:
                        seen[key].append(a)
                        existing_lower.add(a.lower())

    all_concepts = [Concept(name=canonical_by_lower[k], aliases=seen[k]) for k in order][:8]

    seen_topics: set[str] = set()
    all_topics: list[str] = []
    for r in results:
        for t in r.suggested_topics:
            if t.lower() not in seen_topics:
                seen_topics.add(t.lower())
                all_topics.append(t)

    min_result = min(results, key=lambda r: _QUALITY_RANK.get(r.quality, 1))

    merged_language = next((r.language for r in results if r.language), None)

    return AnalysisResult(
        summary=results[0].summary,
        concepts=all_concepts,
        suggested_topics=all_topics[:5],
        quality=min_result.quality,
        language=merged_language,
    )


def _analyze_body(
    body: str,
    existing_concepts: list[str],
    path_name: str,
    client: LLMClientProtocol,
    config: Config,
) -> AnalysisResult:
    """Analyze note body, splitting into chunks when body exceeds configured chunk size."""
    ratio = max(0.25, min(config.pipeline.ingest_chunk_ratio, 0.9))
    chunk_size = max(1, int(config.effective_provider.fast_ctx * ratio))

    if len(body) <= chunk_size:
        prompt = _build_analysis_prompt(body, existing_concepts, path_name)
        return request_structured(
            client=client,
            prompt=prompt,
            model_class=AnalysisResult,
            model=config.models.fast,
            system=_SYSTEM,
            num_ctx=config.effective_provider.fast_ctx,
            max_retries=config.pipeline.ingest_max_retries,
            telemetry_config=config,
            telemetry_stage="ingest_analysis",
        )

    # Split into chunks — no overlap needed for concept extraction
    chunks = [body[i : i + chunk_size] for i in range(0, len(body), chunk_size)]
    log.info(
        "Note %s split into %d chunks for analysis (%d chars, chunk_size=%d)",
        path_name or "unknown",
        len(chunks),
        len(body),
        chunk_size,
    )

    def _analyze_chunk(chunk: str, idx: int) -> AnalysisResult:
        label = f"[part {idx + 1}/{len(chunks)}]"
        log.info("Analyzing %s %s …", path_name or "note", label)
        t0 = time.monotonic()
        prompt = _build_analysis_prompt(chunk, existing_concepts, path_name, chunk_label=label)
        result = request_structured(
            client=client,
            prompt=prompt,
            model_class=AnalysisResult,
            model=config.models.fast,
            system=_SYSTEM,
            num_ctx=config.effective_provider.fast_ctx,
            max_retries=config.pipeline.ingest_max_retries,
            telemetry_config=config,
            telemetry_stage="ingest_analysis_chunk",
        )
        log.info("Analyzed %s %s (%.1fs)", path_name or "note", label, time.monotonic() - t0)
        return result

    if config.pipeline.ingest_parallel:
        from concurrent.futures import ThreadPoolExecutor, as_completed

        env_parallel = os.getenv("OLLAMA_NUM_PARALLEL", "").strip()
        try:
            configured_workers = int(env_parallel) if env_parallel else 0
        except ValueError:
            configured_workers = 0
        if configured_workers <= 0:
            configured_workers = 4

        max_workers = max(1, min(len(chunks), configured_workers))
        chunk_results: list[AnalysisResult | None] = [None] * len(chunks)
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {
                executor.submit(_analyze_chunk, chunk, i): i for i, chunk in enumerate(chunks)
            }
            for future in as_completed(futures):
                chunk_results[futures[future]] = future.result()
        results = [r for r in chunk_results if r is not None]
    else:
        results = [_analyze_chunk(chunk, i) for i, chunk in enumerate(chunks)]

    return _merge_chunk_results(results)


_STOPWORDS = frozenset(
    {
        "a",
        "an",
        "the",
        "and",
        "or",
        "but",
        "is",
        "it",
        "in",
        "on",
        "at",
        "to",
        "by",
        "for",
        "of",
        "as",
        "from",
        "with",
        "this",
        "that",
        "these",
        "those",
        "be",
        "are",
    }
)

_CONCEPT_JUNK_RE = re.compile(
    r"^(?:page|p\.?|figure|fig|section|sec|chapter|ch|slide)\s*[-:#.]?\s*\d+(?:\.\d+)*\b",
    re.IGNORECASE,
)

# Bare dates: "2024-03-15", "3/15/2024", "15.03.2024", "January 2024", "Jan 2025"
_MONTHS_INNER = (
    r"january|february|march|april|may|june|july|august|september|october|november|december"
    r"|jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec"
)
_CONCEPT_TEMPORAL_RE = re.compile(
    r"^(?:"
    r"\d{4}-\d{2}-\d{2}"  # ISO: 2024-03-15
    r"|\d{1,2}/\d{1,2}/\d{2,4}"  # US: 3/15/2024
    r"|\d{1,2}\.\d{2}\.\d{4}"  # EU: 15.03.2024
    r"|(?:" + _MONTHS_INNER + r")\s+\d{4}"  # Named month: January 2024
    r"|\d{4}\s+(?:" + _MONTHS_INNER + r")"  # Year-first: 2024 January
    r")$",
    re.IGNORECASE,
)

# Meeting/transcript timestamps: "00:05:23", "9:15", "10:30 AM", "9:15:00 PM"
_CONCEPT_TIMESTAMP_RE = re.compile(
    r"^\d{1,2}:\d{2}(?::\d{2})?(?:\s*[AP]M)?$",
    re.IGNORECASE,
)

_SINGULAR_EXCEPTIONS = frozenset({"analysis", "news", "series", "species", "status"})


def _is_valid_concept_name(name: str) -> bool:
    """Drop obviously low-information concepts like page/figure labels, bare dates, timestamps, and index entries."""
    stripped = name.strip()
    if not stripped:
        return False
    if _CONCEPT_JUNK_RE.match(stripped):
        return False
    if re.fullmatch(r"\d+(?:\.\d+)*", stripped):
        return False
    # Reject bare dates and timestamps — auto-generated context markers from transcripts/PDFs
    if _CONCEPT_TEMPORAL_RE.fullmatch(stripped):
        return False
    if _CONCEPT_TIMESTAMP_RE.fullmatch(stripped):
        return False
    # Reject standalone "index" — navigation artifact, not a knowledge concept
    if stripped.lower() == "index":
        return False
    return True


def _singularize_token(token: str) -> str:
    """Apply lightweight plural reduction for simple concept deduping."""
    lower = token.lower()
    if len(lower) <= 3 or lower in _SINGULAR_EXCEPTIONS:
        return token
    if lower.endswith("ies") and len(lower) > 4:
        return token[:-3] + "y"
    if lower.endswith("ses") and len(lower) > 4:
        return token[:-2]
    if lower.endswith("s") and not lower.endswith(("ss", "us", "is")):
        return token[:-1]
    return token


def _singularize_phrase(name: str) -> str:
    words = name.split()
    if not words:
        return name
    words[-1] = _singularize_token(words[-1])
    return " ".join(words)


def _concept_singular_key(name: str) -> str:
    return re.sub(r"\s+", " ", _singularize_phrase(name).strip()).lower()


def _validate_aliases(canonical: str, raw_aliases: list[str]) -> list[str]:
    """Filter LLM-produced aliases: remove too-short, stopwords, self-matches, duplicates."""
    seen = {canonical.lower()}
    valid: list[str] = []
    for alias in raw_aliases:
        a = alias.strip()
        if not a or a.lower() in seen:
            continue
        if len(a) < 2:
            continue
        if len(a) <= 3 and not a.isupper():
            continue
        if a.lower() in _STOPWORDS:
            continue
        seen.add(a.lower())
        valid.append(a)
    return valid[:5]


def _normalize_concepts(raw_concepts: list[Concept], db: StateDB) -> list[tuple[str, list[str]]]:
    """Case-insensitive dedup against existing canonical concept names.

    Returns (canonical_name, validated_aliases) pairs.
    """
    existing_names = db.list_all_concept_names()
    existing = {n.lower(): n for n in existing_names}
    existing_singular = {_concept_singular_key(n): n for n in existing_names}
    seen: set[str] = set()
    result: list[tuple[str, list[str]]] = []
    for concept in raw_concepts:
        name = concept.name.strip()
        if not _is_valid_concept_name(name):
            continue
        canonical = existing.get(name.lower()) or existing_singular.get(_concept_singular_key(name))
        if canonical is None:
            canonical = name

        canonical_key = _concept_singular_key(canonical)
        if canonical_key in seen:
            continue
        seen.add(canonical_key)
        aliases = _validate_aliases(canonical, concept.aliases)
        result.append((canonical, aliases))
    return result


_HEADER_SCAN_LINES = 30  # only strip short lines from the opening section

# Media reference patterns for source page preservation
_OBSIDIAN_EMBED_RE = re.compile(
    r"!\[\[([^\]]+\.(?:png|jpg|jpeg|gif|svg|webp|bmp|tiff|avif|pdf|mp4|webm|mov|mp3|wav|ogg))\]\]",
    re.IGNORECASE,
)
_MD_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
_PAGE_FILE_RE = re.compile(r"^page-(\d+)\.md$", re.IGNORECASE)


def _preprocess_web_clip(content: str) -> str:
    """Clean common Obsidian Web Clipper artifacts (nav bars, cookie banners, HTML tags).

    HTML stripping is scoped to the first _HEADER_SCAN_LINES only — body HTML
    (<details>, <kbd>, <sup>, etc.) is intentional and preserved.
    """
    _MD_STARTS = ("#", "-", "*", ">", "[", "!")  # markdown structural chars — always keep
    lines = content.splitlines()

    cleaned = []
    for i, line in enumerate(lines):
        if i < _HEADER_SCAN_LINES:
            # Strip HTML only in header region (nav/banner cleanup)
            line = re.sub(r"<[^>]+>", "", line)
            stripped = line.strip()
            # Skip short non-empty non-markdown lines (nav/banner heuristic)
            if stripped and len(stripped.split()) <= 5 and not stripped.startswith(_MD_STARTS):
                continue
        cleaned.append(line)
    return "\n".join(cleaned)


def _extract_page_number(path: Path) -> int | None:
    m = _PAGE_FILE_RE.match(path.name)
    if not m:
        return None
    try:
        return int(m.group(1))
    except ValueError:
        return None


def _is_ingest_candidate(path: Path) -> bool:
    return path.is_file() and "processed" not in path.parts and not path.name.startswith(".")


def _conversion_output_dir(source_path: Path, config: Config | None) -> Path:
    """Return the directory under conversions/ that mirrors the source file's location in raw/."""
    if config is not None:
        try:
            rel = source_path.relative_to(config.raw_dir)
        except ValueError:
            rel = Path(source_path.stem)
        return config.conversions_dir / rel.parent / sanitize_filename(source_path.stem)
    # fallback when no config available (e.g. tests calling converters directly)
    return source_path.parent / sanitize_filename(source_path.stem)


def _cleanup_source_summary_mirror(
    output_dir: Path,
    config: Config,
) -> None:
    """Remove stale mirrored source summaries for a reconverted conversion folder."""
    try:
        rel_output = output_dir.relative_to(config.raw_dir)
    except ValueError:
        try:
            rel_output = output_dir.relative_to(config.conversions_dir)
        except ValueError:
            return

    source_dir = _source_summary_dir(config, rel_output)
    if not source_dir.exists():
        return

    for pattern in ("page-*.md", "group-*.md"):
        for existing in source_dir.glob(pattern):
            existing.unlink(missing_ok=True)


def _cleanup_page_source_summaries(config: Config, page_paths: list[Path]) -> None:
    """Remove stale wiki/sources mirrors for migrated legacy page-*.md files."""
    for page_path in page_paths:
        try:
            source_path = _source_summary_path(config, page_path)
        except Exception:
            continue
        source_path.unlink(missing_ok=True)


def _write_pdf_group_source_mirror(group_path: Path, config: Config) -> None:
    """Write a deterministic source mirror for grouped PDF markdown.

    This keeps wiki/sources aligned with grouped raw outputs immediately after
    conversion, even before a full LLM ingest refresh runs.
    """
    try:
        src_meta, body = parse_note(group_path)
    except Exception:
        src_meta, body = {}, _read_text_with_fallback(group_path)

    rel_raw = group_path.relative_to(config.vault).as_posix()
    out_path = _source_summary_path(config, group_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    title = src_meta.get("title") or group_path.stem.replace("-", " ").title()
    first_para = next((line.strip() for line in body.splitlines() if line.strip()), "")
    summary = first_para[:400] if first_para else "Grouped PDF extract."
    source_url = src_meta.get("source") or src_meta.get("url") or ""
    now = datetime.now().strftime("%Y-%m-%d")

    out_meta: dict = {
        "title": title,
        "aliases": generate_aliases(title, ""),
        "tags": ["source", "pdf-group"],
        "status": "published",
        "source_file": rel_raw,
        "quality": "medium",
        "created": now,
    }
    if source_url:
        out_meta["source_url"] = source_url

    body_parts = [
        f"# {title}",
        "",
        "## Summary",
        summary,
        "",
        "## Source Info",
        "- **Quality:** medium",
        f"- **Raw file:** [[{rel_raw}]]",
        f"- **Ingested:** {now}",
    ]
    if source_url:
        body_parts.append(f"- **URL:** {source_url}")

    write_note(out_path, out_meta, "\n".join(body_parts))


def _should_rebuild_pdf_groups(pdf_path: Path, output_dir: Path) -> bool:
    """Decide whether grouped markdown needs to be rebuilt for a PDF.

    Rebuild when grouped outputs are missing, when legacy page files are present,
    or when the source PDF is newer than current grouped outputs.
    """
    existing_groups = sorted(output_dir.glob("group-*.md")) if output_dir.exists() else []
    if not existing_groups:
        return True

    if any(output_dir.glob("page-*.md")):
        return True

    try:
        pdf_mtime = pdf_path.stat().st_mtime
    except OSError:
        return False

    newest_group_mtime = max((p.stat().st_mtime for p in existing_groups), default=0.0)
    return pdf_mtime > newest_group_mtime


def _should_rebuild_converted(source_path: Path, out_path: Path) -> bool:
    """Rebuild the converted markdown when the output is missing or the source is newer."""
    if not out_path.exists():
        return True
    try:
        return source_path.stat().st_mtime > out_path.stat().st_mtime
    except OSError:
        return False


def _is_section_boundary(text: str, patterns: list[str]) -> bool:
    first_non_empty = ""
    for line in text.splitlines():
        stripped = line.strip()
        if stripped:
            first_non_empty = stripped
            break
    if not first_non_empty:
        return False
    lowered = first_non_empty.lower()
    for pattern in patterns:
        if re.match(pattern, lowered, flags=re.IGNORECASE):
            return True
    return False


def _migrate_pagewise_markdown_dir(dir_path: Path, config: Config) -> list[Path]:
    """Convert legacy page-*.md files in a folder into grouped markdown files.

    Migration is idempotent: it rewrites grouped files from page files, then
    removes legacy pages and their source-summary mirrors.
    """
    page_paths = [
        p for p in sorted(dir_path.glob("page-*.md")) if p.is_file() and _extract_page_number(p) is not None
    ]
    if not page_paths:
        return []

    pipeline = config.pipeline
    preserve_markers = pipeline.pdf_preserve_page_markers
    strategy = pipeline.pdf_split_strategy
    max_pages = pipeline.pdf_max_chunk_pages
    min_chars = pipeline.pdf_min_chunk_chars
    max_chars = pipeline.pdf_max_chunk_chars
    section_patterns = pipeline.pdf_section_patterns

    pages: list[tuple[int, str]] = []
    for page_path in page_paths:
        page_number = _extract_page_number(page_path)
        if page_number is None:
            continue
        try:
            _, page_body = parse_note(page_path)
        except Exception:
            page_body = _read_text_with_fallback(page_path)
        pages.append((page_number, page_body.strip()))

    if not pages:
        return []

    pages.sort(key=lambda item: item[0])
    if strategy == "per-page":
        groups = [[item] for item in pages]
    else:
        groups = _group_pdf_pages(
            pages,
            max_pages=max_pages,
            min_chars=min_chars,
            max_chars=max_chars,
            section_patterns=section_patterns,
        )

    written_paths: list[Path] = []
    source_label = dir_path.name.replace("_", " ").strip() or "Grouped Pages"

    for group in groups:
        start_page = group[0][0]
        end_page = group[-1][0]
        out_path = dir_path / f"group-{start_page:03d}-{end_page:03d}.md"

        blocks: list[str] = []
        for page_number, text in group:
            if preserve_markers:
                if re.match(r"^\[\s*page\s+\d+\s*\]", text.strip(), flags=re.IGNORECASE):
                    blocks.append(text)
                else:
                    blocks.append(f"[Page {page_number}]\n\n{text}")
            else:
                blocks.append(text)

        body = "\n\n---\n\n".join(blocks)
        write_note(
            out_path,
            {
                "title": f"{source_label} - Pages {start_page}-{end_page}",
                "source_pages": [page for page, _ in group],
                "source_page_start": start_page,
                "source_page_end": end_page,
                "tags": ["pdf-group"],
            },
            f"## Pages {start_page}-{end_page}\n\n{body}\n",
        )
        written_paths.append(out_path)

    _cleanup_page_source_summaries(config, page_paths)
    for page_path in page_paths:
        page_path.unlink(missing_ok=True)

    return sorted(written_paths)


def _group_pdf_pages(
    pages: list[tuple[int, str]],
    max_pages: int,
    min_chars: int,
    max_chars: int,
    section_patterns: list[str],
) -> list[list[tuple[int, str]]]:
    groups: list[list[tuple[int, str]]] = []
    current_group: list[tuple[int, str]] = []
    current_chars = 0

    for page_number, text in pages:
        page_chars = len(text)
        starts_new_section = _is_section_boundary(text, section_patterns)

        should_flush = bool(current_group) and (
            len(current_group) >= max_pages
            or current_chars + page_chars > max_chars
            or (starts_new_section and current_chars >= min_chars)
        )

        if should_flush:
            groups.append(current_group)
            current_group = []
            current_chars = 0

        current_group.append((page_number, text))
        current_chars += page_chars

    if current_group:
        groups.append(current_group)

    return groups


def convert_pdf_to_markdown(
    pdf_path: Path,
    overwrite: bool = False,
    config: Config | None = None,
) -> list[Path]:
    """Convert a PDF into deterministic grouped markdown files inside a sibling folder."""
    try:
        from pypdf import PdfReader
    except Exception as e:
        log.warning("PDF conversion unavailable for %s: %s", pdf_path.name, e)
        return []

    output_dir = _conversion_output_dir(pdf_path, config)
    existing_groups = sorted(output_dir.glob("group-*.md")) if output_dir.exists() else []
    if existing_groups and not overwrite:
        return existing_groups

    try:
        reader = PdfReader(str(pdf_path))
    except Exception as e:
        log.warning("Failed to read PDF %s: %s", pdf_path.name, e)
        return []

    output_dir.mkdir(parents=True, exist_ok=True)
    if overwrite:
        for existing in output_dir.glob("page-*.md"):
            existing.unlink(missing_ok=True)
        for existing in output_dir.glob("group-*.md"):
            existing.unlink(missing_ok=True)
        if config is not None:
            _cleanup_source_summary_mirror(output_dir, config)

    pipeline = config.pipeline if config else None
    strategy = pipeline.pdf_split_strategy if pipeline else "grouped"
    max_pages = pipeline.pdf_max_chunk_pages if pipeline else 4
    min_chars = pipeline.pdf_min_chunk_chars if pipeline else 800
    max_chars = pipeline.pdf_max_chunk_chars if pipeline else 14000
    section_patterns = pipeline.pdf_section_patterns if pipeline else [r"^#", r"^chapter\\b"]
    preserve_markers = pipeline.pdf_preserve_page_markers if pipeline else True

    vision_model = (config.pipeline.vision_model if config and config.pipeline else "").strip()
    vision_client = None
    if vision_model:
        from ..ollama_client import OllamaClient

        provider_url = config.effective_provider.url if config else "http://localhost:11434"
        vision_client = OllamaClient(base_url=provider_url)

    rel_pdf = pdf_path.as_posix()
    extracted_pages: list[tuple[int, str]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()

        image_descriptions: list[str] = []
        if vision_client and vision_model:
            for img in getattr(page, "images", []):
                try:
                    b64 = base64.b64encode(img.data).decode("ascii")
                    desc = vision_client.generate(
                        prompt="Describe this image concisely in 1-2 sentences for a knowledge base. Focus on content and key information.",
                        model=vision_model,
                        images=[b64],
                    )
                    if desc.strip():
                        image_descriptions.append(f"> **Image:** {desc.strip()}")
                except Exception as exc:
                    log.debug("Vision LLM failed for image in %s p%d: %s", pdf_path.name, page_number, exc)

        if not text and not image_descriptions:
            log.debug("Skipping empty page %d in %s", page_number, pdf_path.name)
            continue

        if image_descriptions:
            text = (text + "\n\n" if text else "") + "\n\n".join(image_descriptions)

        extracted_pages.append((page_number, text))

    if not extracted_pages:
        return []

    groups: list[list[tuple[int, str]]]
    if strategy == "per-page":
        groups = [[item] for item in extracted_pages]
    else:
        groups = _group_pdf_pages(
            extracted_pages,
            max_pages=max_pages,
            min_chars=min_chars,
            max_chars=max_chars,
            section_patterns=section_patterns,
        )

    written_paths: list[Path] = []
    for group in groups:
        start_page = group[0][0]
        end_page = group[-1][0]
        out_path = output_dir / f"group-{start_page:03d}-{end_page:03d}.md"

        if preserve_markers:
            content_blocks = [f"[Page {page}]\n\n{text}" for page, text in group]
        else:
            content_blocks = [text for _, text in group]
        body = "\n\n---\n\n".join(content_blocks)

        write_note(
            out_path,
            {
                "title": f"{pdf_path.stem} - Pages {start_page}-{end_page}",
                "source_pdf": rel_pdf,
                "source_pages": [page for page, _ in group],
                "source_page_start": start_page,
                "source_page_end": end_page,
                "tags": ["pdf-group"],
            },
            f"## Pages {start_page}-{end_page}\n\n{body}\n",
        )
        if config is not None:
            _write_pdf_group_source_mirror(out_path, config)
        written_paths.append(out_path)

    log.info(
        "Converted PDF %s into %d grouped markdown file(s) under %s",
        pdf_path.name,
        len(written_paths),
        output_dir.name,
    )
    emit_event(
        config=None,
        event_type="pdf_converted",
        source_pdf=pdf_path.name,
        page_count=len(written_paths),
    )
    return written_paths


def convert_xlsx_to_markdown(
    xlsx_path: Path,
    overwrite: bool = False,
    config: Config | None = None,
) -> list[Path]:
    """Convert an Excel workbook to markdown. Each sheet becomes a ## section with a table."""
    try:
        import openpyxl
    except ImportError:
        log.warning("openpyxl not installed; skipping %s. Install: uv add openpyxl", xlsx_path.name)
        return []

    out_dir = _conversion_output_dir(xlsx_path, config)
    out_path = out_dir / "converted.md"
    if out_path.exists() and not overwrite and not _should_rebuild_converted(xlsx_path, out_path):
        return [out_path]

    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    except Exception as e:
        log.warning("Failed to read Excel %s: %s", xlsx_path.name, e)
        return []

    sections: list[str] = []
    for sheet in wb.worksheets:
        rows = [r for r in sheet.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        header = [str(c) if c is not None else "" for c in rows[0]]
        col_count = len(header)
        lines = [f"## {sheet.title}", "", "| " + " | ".join(header) + " |", "|" + " --- |" * col_count]
        for row in rows[1:]:
            cells = [(str(c) if c is not None else "") for c in row]
            padded = (cells + [""] * col_count)[:col_count]
            lines.append("| " + " | ".join(padded) + " |")
        sections.append("\n".join(lines))

    if not sections:
        return []

    body = "\n\n".join(sections) + "\n"
    out_dir.mkdir(parents=True, exist_ok=True)
    write_note(
        out_path,
        {"title": xlsx_path.stem, "source_xlsx": xlsx_path.as_posix(), "tags": ["xlsx-converted"]},
        body,
    )
    log.info("Converted Excel %s → %s (%d sheet(s))", xlsx_path.name, out_path.name, len(sections))
    return [out_path]


def convert_docx_to_markdown(
    docx_path: Path,
    overwrite: bool = False,
    config: Config | None = None,
) -> list[Path]:
    """Convert a Word document to markdown, preserving heading hierarchy and tables."""
    try:
        from docx import Document
    except ImportError:
        log.warning("python-docx not installed; skipping %s. Install: uv add python-docx", docx_path.name)
        return []

    out_dir = _conversion_output_dir(docx_path, config)
    out_path = out_dir / "converted.md"
    if out_path.exists() and not overwrite and not _should_rebuild_converted(docx_path, out_path):
        return [out_path]

    try:
        doc = Document(str(docx_path))
    except Exception as e:
        log.warning("Failed to read Word doc %s: %s", docx_path.name, e)
        return []

    _HEADING_PREFIXES = {f"heading {i}": "#" * i for i in range(1, 5)}
    lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name or "").lower() if para.style else ""
        prefix = _HEADING_PREFIXES.get(style_name, "")
        lines.append(f"{prefix} {text}" if prefix else text)

    for i, table in enumerate(doc.tables):
        rows = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in table.rows]
        if not rows:
            continue
        col_count = len(rows[0])
        lines.append(f"\n## Table {i + 1}\n")
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("|" + " --- |" * col_count)
        for row in rows[1:]:
            padded = (row + [""] * col_count)[:col_count]
            lines.append("| " + " | ".join(padded) + " |")

    if not lines:
        return []

    body = "\n".join(lines) + "\n"
    out_dir.mkdir(parents=True, exist_ok=True)
    write_note(
        out_path,
        {"title": docx_path.stem, "source_docx": docx_path.as_posix(), "tags": ["docx-converted"]},
        body,
    )
    log.info("Converted Word doc %s → %s", docx_path.name, out_path.name)
    return [out_path]


def convert_pptx_to_markdown(
    pptx_path: Path,
    overwrite: bool = False,
    config: Config | None = None,
) -> list[Path]:
    """Convert a PowerPoint presentation to markdown. Each slide becomes a ## section."""
    try:
        from pptx import Presentation
    except ImportError:
        log.warning("python-pptx not installed; skipping %s. Install: uv add python-pptx", pptx_path.name)
        return []

    out_dir = _conversion_output_dir(pptx_path, config)
    out_path = out_dir / "converted.md"
    if out_path.exists() and not overwrite and not _should_rebuild_converted(pptx_path, out_path):
        return [out_path]

    try:
        prs = Presentation(str(pptx_path))
    except Exception as e:
        log.warning("Failed to read PowerPoint %s: %s", pptx_path.name, e)
        return []

    sections: list[str] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_text = (title_shape.text or "").strip() if title_shape else ""
        heading = f"## Slide {slide_num}: {title_text}" if title_text else f"## Slide {slide_num}"

        body_lines: list[str] = []
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            if shape is title_shape:
                continue
            text = shape.text_frame.text.strip()
            if text:
                body_lines.append(text)

        parts = [heading]
        if body_lines:
            parts.append("\n".join(body_lines))
        sections.append("\n\n".join(parts))

    if not sections:
        return []

    body = "\n\n".join(sections) + "\n"
    out_dir.mkdir(parents=True, exist_ok=True)
    write_note(
        out_path,
        {"title": pptx_path.stem, "source_pptx": pptx_path.as_posix(), "tags": ["pptx-converted"]},
        body,
    )
    log.info("Converted PowerPoint %s → %s (%d slide(s))", pptx_path.name, out_path.name, len(sections))
    return [out_path]


def convert_csv_to_markdown(
    csv_path: Path,
    overwrite: bool = False,
    config: Config | None = None,
) -> list[Path]:
    """Convert a CSV file to a markdown table."""
    out_dir = _conversion_output_dir(csv_path, config)
    out_path = out_dir / "converted.md"
    if out_path.exists() and not overwrite and not _should_rebuild_converted(csv_path, out_path):
        return [out_path]

    try:
        with open(csv_path, newline="", encoding="utf-8", errors="replace") as f:
            rows = [r for r in csv.reader(f) if any(c.strip() for c in r)]
    except Exception as e:
        log.warning("Failed to read CSV %s: %s", csv_path.name, e)
        return []

    if not rows:
        return []

    col_count = max(len(r) for r in rows)
    header = (rows[0] + [""] * col_count)[:col_count]
    lines = ["| " + " | ".join(header) + " |", "|" + " --- |" * col_count]
    for row in rows[1:]:
        padded = (row + [""] * col_count)[:col_count]
        lines.append("| " + " | ".join(c.replace("\n", " ") for c in padded) + " |")

    body = "\n".join(lines) + "\n"
    out_dir.mkdir(parents=True, exist_ok=True)
    write_note(
        out_path,
        {"title": csv_path.stem, "source_csv": csv_path.as_posix(), "tags": ["csv-converted"]},
        body,
    )
    log.info("Converted CSV %s → %s", csv_path.name, out_path.name)
    return [out_path]


def _cleanup_legacy_raw_conversion_dirs(config: Config) -> None:
    """Remove old conversion subdirectories that were created inside raw/ before the conversions/ bucket existed.

    Looks for directories inside raw/ whose name matches a source file stem and contains only
    group-*.md or converted.md files — the signatures of old-style in-place conversions.
    """
    if not config.raw_dir.exists():
        return
    source_extensions = {".pdf", ".xlsx", ".docx", ".pptx", ".csv"}
    for source_file in config.raw_dir.rglob("*"):
        if source_file.suffix.lower() not in source_extensions:
            continue
        legacy_dir = source_file.parent / sanitize_filename(source_file.stem)
        if not legacy_dir.is_dir():
            continue
        contents = list(legacy_dir.iterdir())
        if not contents:
            continue
        is_conversion_dir = all(
            f.is_file() and (f.name == "converted.md" or f.name.startswith("group-") or f.name.startswith("page-"))
            for f in contents
        )
        if is_conversion_dir:
            for f in contents:
                f.unlink(missing_ok=True)
            try:
                legacy_dir.rmdir()
                log.info("Removed legacy raw/ conversion dir: %s", legacy_dir.relative_to(config.vault))
            except OSError:
                pass


def collect_ingest_paths(config: Config, paths: list[Path] | None = None) -> list[Path]:
    """Collect markdown paths for ingest, auto-converting PDFs into grouped notes, and allowing arbitrary .md files as first-class input.

    All .md files (not just grouped PDF markdown) are included for chunking, analysis, and summary/concept extraction.
    """
    if paths is None:
        _cleanup_legacy_raw_conversion_dirs(config)
        candidates = list(config.raw_dir.rglob("*")) if config.raw_dir.exists() else []
    else:
        candidates = [Path(path) for path in paths]

    md_paths: list[Path] = []
    seen: set[str] = set()
    migrated_dirs: set[str] = set()

    for path in sorted(candidates):
        if not _is_ingest_candidate(path):
            continue

        suffix = path.suffix.lower()
        if suffix == ".pdf":
            output_dir = _conversion_output_dir(path, config)
            overwrite = _should_rebuild_pdf_groups(path, output_dir)
            for page_path in convert_pdf_to_markdown(path, overwrite=overwrite, config=config):
                key = page_path.resolve().as_posix()
                if key not in seen:
                    seen.add(key)
                    md_paths.append(page_path)
        elif suffix == ".xlsx":
            out_path = _conversion_output_dir(path, config) / "converted.md"
            overwrite = _should_rebuild_converted(path, out_path)
            for converted in convert_xlsx_to_markdown(path, overwrite=overwrite, config=config):
                key = converted.resolve().as_posix()
                if key not in seen:
                    seen.add(key)
                    md_paths.append(converted)
        elif suffix == ".docx":
            out_path = _conversion_output_dir(path, config) / "converted.md"
            overwrite = _should_rebuild_converted(path, out_path)
            for converted in convert_docx_to_markdown(path, overwrite=overwrite, config=config):
                key = converted.resolve().as_posix()
                if key not in seen:
                    seen.add(key)
                    md_paths.append(converted)
        elif suffix == ".pptx":
            out_path = _conversion_output_dir(path, config) / "converted.md"
            overwrite = _should_rebuild_converted(path, out_path)
            for converted in convert_pptx_to_markdown(path, overwrite=overwrite, config=config):
                key = converted.resolve().as_posix()
                if key not in seen:
                    seen.add(key)
                    md_paths.append(converted)
        elif suffix == ".csv":
            out_path = _conversion_output_dir(path, config) / "converted.md"
            overwrite = _should_rebuild_converted(path, out_path)
            for converted in convert_csv_to_markdown(path, overwrite=overwrite, config=config):
                key = converted.resolve().as_posix()
                if key not in seen:
                    seen.add(key)
                    md_paths.append(converted)
        elif suffix == ".md":
            page_number = _extract_page_number(path)
            if page_number is not None:
                dir_key = path.parent.resolve().as_posix()
                if dir_key not in migrated_dirs:
                    migrated_dirs.add(dir_key)
                    migrated = _migrate_pagewise_markdown_dir(path.parent, config)
                    for migrated_path in migrated:
                        key = migrated_path.resolve().as_posix()
                        if key not in seen:
                            seen.add(key)
                            md_paths.append(migrated_path)
                continue
            # Accept all other .md files (not just grouped PDF markdown)
            key = path.resolve().as_posix()
            if key not in seen:
                seen.add(key)
                md_paths.append(path)

    return sorted(md_paths)


def _collect_media_refs(body: str) -> list[str]:
    """Extract media references from note body for preservation in source pages."""
    refs: list[str] = []
    for m in _OBSIDIAN_EMBED_RE.finditer(body):
        refs.append(f"- ![[{m.group(1)}]]")
    for m in _MD_IMAGE_RE.finditer(body):
        alt, url = m.group(1), m.group(2)
        refs.append(f"- ![{alt}]({url})")
    return refs


def _extract_existing_source_concepts(body: str) -> list[str]:
    """Extract concept link targets from an existing source summary Concepts section."""
    lines = body.splitlines()
    in_concepts = False
    collected: list[str] = []
    seen: set[str] = set()

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("## "):
            in_concepts = stripped.lower() == "## concepts"
            continue
        if not in_concepts or not stripped:
            continue
        if not stripped.startswith("-"):
            continue

        for match in re.finditer(r"\[\[([^\]|]+)(?:\|[^\]]+)?\]\]", stripped):
            target = sanitize_wikilink_target(match.group(1).strip())
            if not target:
                continue
            if not _is_valid_concept_name(target):
                continue
            key = target.lower()
            if key in seen:
                continue
            seen.add(key)
            collected.append(target)

    return collected


def _create_source_summary_page(
    path: Path,
    src_meta: dict,
    result: AnalysisResult,
    config: Config,
    body: str = "",
) -> Path:
    """
    Generate wiki/sources/{Title}.md from AnalysisResult. No extra LLM call.
    Returns the path written.
    """
    # Derive title from note frontmatter > file stem
    title = src_meta.get("title") or path.stem.replace("-", " ").title()
    out_path = _source_summary_path(config, path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    now = datetime.now().strftime("%Y-%m-%d")
    rel_raw = path.relative_to(config.vault).as_posix()
    source_url = src_meta.get("source") or src_meta.get("url") or ""
    aliases = generate_aliases(title, "")  # source pages rarely have abbreviations

    existing_concepts: list[str] = []
    if out_path.exists():
        try:
            _, existing_body = parse_note(out_path)
            existing_concepts = _extract_existing_source_concepts(existing_body)
        except Exception:
            existing_concepts = []

    merged_concepts: list[str] = []
    concept_seen: set[str] = set()
    for concept in existing_concepts:
        key = concept.lower()
        if key in concept_seen:
            continue
        concept_seen.add(key)
        merged_concepts.append(concept)
    for concept in result.concepts[:8]:
        if not _is_valid_concept_name(concept.name):
            continue
        target = sanitize_wikilink_target(concept.name)
        if not target:
            continue
        key = target.lower()
        if key in concept_seen:
            continue
        concept_seen.add(key)
        merged_concepts.append(target)

    concept_lines = "\n".join(f"- [[{target}]]" for target in merged_concepts)

    out_meta: dict = {
        "title": title,
        "aliases": aliases,
        "tags": ["source"],
        "status": "published",
        "source_file": rel_raw,
        "quality": result.quality,
        "created": now,
    }
    if source_url:
        out_meta["source_url"] = source_url

    body_parts = [
        f"# {title}",
        "",
        "## Summary",
        result.summary,
        "",
        "## Concepts",
        concept_lines,
        "",
        "## Source Info",
        f"- **Quality:** {result.quality}",
        f"- **Raw file:** [[{rel_raw}]]",
        f"- **Ingested:** {now}",
    ]
    if source_url:
        body_parts.append(f"- **URL:** {source_url}")

    media_refs = _collect_media_refs(body)
    if media_refs:
        body_parts += ["", "## Media"] + media_refs

    write_note(out_path, out_meta, "\n".join(body_parts))
    log.info("Source summary written: %s", out_path.name)
    return out_path


def _ensure_source_summary_for_existing_ingest(
    path: Path,
    src_meta: dict,
    body: str,
    record: RawNoteRecord,
    db: StateDB,
    config: Config,
) -> None:
    """Backfill a missing source summary for already-ingested notes without re-analysis."""
    out_path = _source_summary_path(config, path)
    if out_path.exists():
        return

    rel_path = path.relative_to(config.vault).as_posix()
    concept_names = db.get_concepts_for_sources([rel_path])[:8]
    concepts = [Concept(name=name, aliases=[]) for name in concept_names]
    quality = record.quality if record.quality in {"low", "medium", "high"} else "medium"
    summary = (record.summary or "").strip() or "No summary available."
    result = AnalysisResult(
        summary=summary,
        concepts=concepts,
        suggested_topics=[],
        quality=quality,
        language=record.language,
    )
    _create_source_summary_page(path, src_meta, result, config, body=body)
    log.info("Rebuilt missing source summary for already-ingested note: %s", path.name)


def _find_document_source_dirs(config: Config) -> list[Path]:
    """Return wiki/sources/ subdirectories that contain group-*.md files (PDF documents)."""
    if not config.sources_dir.exists():
        return []
    return sorted(d for d in config.sources_dir.rglob("*") if d.is_dir() and any(d.glob("group-*.md")))


def _write_document_aggregate(source_dir: Path, config: Config, db: StateDB) -> Path | None:
    """
    Write or refresh the document-level aggregation file for a PDF document.

    source_dir is a wiki/sources/<document-name>/ directory.
    Aggregation file lands at source_dir.parent/<document-name>.md.
    Returns the path written, or None if already up-to-date.
    """
    group_files = sorted(source_dir.glob("group-*.md"))
    if not group_files:
        return None

    # Parse each group's source_file reference and look up DB records
    groups: list[dict] = []
    for gf in group_files:
        try:
            meta, _ = parse_note(gf)
        except Exception:
            meta = {}
        source_file: str = meta.get("source_file", "") if isinstance(meta, dict) else ""
        record = db.get_raw(source_file) if source_file else None
        concepts = db.get_concepts_for_sources([source_file]) if source_file else []
        groups.append(
            {
                "file": gf,
                "meta": meta,
                "source_file": source_file,
                "record": record,
                "concepts": concepts,
            }
        )

    # Compute incremental fingerprint from content hashes stored in the DB
    hash_inputs: list[str] = []
    for g in groups:
        r = g["record"]
        if r and r.content_hash:
            hash_inputs.append(r.content_hash)
        else:
            try:
                hash_inputs.append(hashlib.sha256(g["file"].read_bytes()).hexdigest()[:16])
            except OSError:
                hash_inputs.append(g["source_file"] or g["file"].name)

    current_sig = hashlib.sha256("\n".join(sorted(hash_inputs)).encode()).hexdigest()[:16]

    # Skip regeneration when nothing changed
    agg_path = source_dir.parent / (source_dir.name + ".md")
    if agg_path.exists():
        try:
            existing_meta, _ = parse_note(agg_path)
            if existing_meta.get("group_sig") == current_sig:
                log.debug("Document aggregate up-to-date: %s", source_dir.name)
                return None
        except Exception:
            pass

    # Merge concepts across all groups (most cross-group concepts first, cap 30)
    concept_counts: dict[str, int] = {}
    concept_canonical: dict[str, str] = {}
    for g in groups:
        for name in g["concepts"]:
            key = name.lower()
            concept_counts[key] = concept_counts.get(key, 0) + 1
            if key not in concept_canonical:
                concept_canonical[key] = name
    merged_concepts = [
        concept_canonical[k] for k in sorted(concept_counts, key=lambda k: -concept_counts[k])
    ][:30]

    # Best summary: highest-quality group with a non-empty DB summary
    best_summary = ""
    best_rank = -1
    for g in groups:
        r = g["record"]
        if r and r.summary and r.summary.strip():
            rank = _QUALITY_RANK.get(r.quality or "medium", 1)
            if rank > best_rank:
                best_rank = rank
                best_summary = r.summary.strip()
    if not best_summary:
        best_summary = f"Aggregated source document with {len(group_files)} group(s)."

    # Overall quality: minimum across recorded qualities (conservative)
    recorded = [g["record"].quality for g in groups if g["record"] and g["record"].quality in _QUALITY_RANK]
    overall_quality = min(recorded, key=lambda q: _QUALITY_RANK[q]) if recorded else "medium"

    # Quality breakdown for display
    q_counts: dict[str, int] = {}
    for g in groups:
        r = g["record"]
        key = r.quality if r and r.quality in _QUALITY_RANK else "failed"
        q_counts[key] = q_counts.get(key, 0) + 1
    q_parts = [f"{q_counts[k]} {k}" for k in ("high", "medium", "low", "failed") if q_counts.get(k)]
    quality_breakdown = " · ".join(q_parts)

    # Detect the source PDF path (raw/<...>/<document-name>.pdf)
    source_pdf = ""
    try:
        rel_to_sources = source_dir.relative_to(config.sources_dir)
        candidate = config.raw_dir / rel_to_sources.parent / (source_dir.name + ".pdf")
        if candidate.exists():
            source_pdf = candidate.relative_to(config.vault).as_posix()
    except (ValueError, OSError):
        pass

    # Page group wikilinks using <document-name>/<group-stem> for disambiguation
    page_group_lines: list[str] = []
    for g in groups:
        gf = g["file"]
        title = g["meta"].get("title", gf.stem) if isinstance(g["meta"], dict) else gf.stem
        page_group_lines.append(f"- [[{source_dir.name}/{gf.stem}|{title}]]")

    now = datetime.now().strftime("%Y-%m-%d")
    doc_title = source_dir.name

    out_meta: dict = {
        "title": doc_title,
        "aliases": [doc_title.lower()],
        "tags": ["source", "source-document"],
        "status": "published",
        "quality": overall_quality,
        "group_count": len(group_files),
        "group_sig": current_sig,
        "created": now,
        "updated": now,
    }
    if source_pdf:
        out_meta["source_pdf"] = source_pdf

    concept_lines = "\n".join(f"- [[{name}]]" for name in merged_concepts)
    page_groups_text = "\n".join(page_group_lines)

    body = "\n".join(
        [
            f"# {doc_title}",
            "",
            "## Summary",
            best_summary,
            "",
            "## Concepts",
            concept_lines,
            "",
            "## Page Groups",
            page_groups_text,
            "",
            "## Quality Rollup",
            f"- **Overall quality:** {overall_quality}",
            f"- **Groups:** {len(group_files)} ({quality_breakdown})",
        ]
    )

    write_note(agg_path, out_meta, body)
    log.info("Document aggregate written: %s", agg_path.name)
    return agg_path


def ingest_note(
    path: Path,
    config: Config,
    client: LLMClientProtocol,
    db: StateDB,
    rag=None,  # Optional RAGStore, injected in Phase 2
    existing_topics: list[str] | None = None,  # existing concept names for prompt
    force: bool = False,
) -> AnalysisResult | None:
    """
    Ingest a single raw note (.md file or grouped PDF markdown).

    Returns AnalysisResult or None if skipped (duplicate / already ingested).
    Handles arbitrary .md files as first-class input, chunking and analyzing them for summary and concepts extraction.
    """
    fn_t0 = time.monotonic()
    try:
        meta, body = parse_note(path)
    except Exception:
        meta, body = {}, _read_text_with_fallback(path)

    # Hash body only (strip frontmatter) so copies are detected as duplicates
    # even after ingest has updated the original's frontmatter (olw_status etc.).
    # Exception: when source_pdf is set (PDF-extracted pages), include it in the
    # hash so pages from *different* PDFs with identical text are not falsely
    # flagged as duplicates.
    source_pdf = meta.get("source_pdf", "")
    body_for_hash = body
    hash_input = (source_pdf + "\x00" + body_for_hash) if source_pdf else body_for_hash
    h = _content_hash(hash_input)

    # Dedup check
    rel_path = path.relative_to(config.vault).as_posix()

    existing = db.get_raw_by_hash(h)
    if existing and existing.path != rel_path:
        log.info("Duplicate of %s, skipping %s", existing.path, path.name)
        emit_event(
            config,
            event_type="function_timing",
            function_name="ingest_note",
            stage="ingest",
            model=config.models.fast,
            success=True,
            outcome="skipped_duplicate",
            duration_ms=round((time.monotonic() - fn_t0) * 1000.0, 2),
            note=path.name,
        )
        return None

    record = db.get_raw(rel_path)

    if record and record.status == "ingested" and not force:
        if record.content_hash == h:
            try:
                _ensure_source_summary_for_existing_ingest(path, meta, body, record, db, config)
            except Exception as e:
                log.warning("Source summary backfill failed for %s: %s", path.name, e)
            log.info("Already ingested: %s", path.name)
            emit_event(
                config,
                event_type="function_timing",
                function_name="ingest_note",
                stage="ingest",
                model=config.models.fast,
                success=True,
                outcome="skipped_already_ingested",
                duration_ms=round((time.monotonic() - fn_t0) * 1000.0, 2),
                note=path.name,
            )
            return None
        log.info("Detected content change, re-ingesting: %s", path.name)

    # Pre-process web clips
    if meta.get("source") or meta.get("url"):  # web clipper adds these
        body = _preprocess_web_clip(body)

    # Skip notes with no usable content (e.g. image-only PDF pages already on disk)
    _EMPTY_BODY_MARKERS = (
        "(No extractable text found on this page. The PDF may be image-only.)",
    )
    stripped_body = body.strip()
    if not stripped_body or stripped_body in _EMPTY_BODY_MARKERS:
        log.info("Skipping empty/image-only note: %s", path.name)
        emit_event(
            config,
            event_type="function_timing",
            function_name="ingest_note",
            stage="ingest",
            model=config.models.fast,
            success=True,
            outcome="skipped_empty",
            duration_ms=round((time.monotonic() - fn_t0) * 1000.0, 2),
            note=path.name,
        )
        return None

    # Chunk + embed only when RAG store is wired in (Phase 2)
    if rag is not None:
        chunks = chunk_text(
            body, chunk_size=config.rag.chunk_size, overlap=config.rag.chunk_overlap
        )
        embeddings = client.embed_batch(chunks, model=config.models.embed)
        rag.add_document(
            doc_id=rel_path,
            chunks=chunks,
            embeddings=embeddings,
            metadata={"source": rel_path, "type": "raw"},
        )

    # LLM analysis — use existing concept names so model can reuse canonical names
    if existing_topics is None:
        existing_topics = db.list_all_concept_names()
    try:
        result: AnalysisResult = _analyze_body(
            body=body,
            existing_concepts=existing_topics,
            path_name=path.name,
            client=client,
            config=config,
        )
    except Exception as e:
        log.error("Analysis failed for %s: %s", path.name, e)
        db.upsert_raw(
            RawNoteRecord(
                path=rel_path,
                content_hash=h,
                status="failed",
                error=str(e),
            )
        )
        emit_event(
            config,
            event_type="function_timing",
            function_name="ingest_note",
            stage="ingest",
            model=config.models.fast,
            success=False,
            outcome="failed",
            duration_ms=round((time.monotonic() - fn_t0) * 1000.0, 2),
            note=path.name,
            error_class=e.__class__.__name__,
            error_message=str(e),
        )
        return None

    # Update state DB (raw files stay immutable — metadata lives in state.db only)
    db.upsert_raw(
        RawNoteRecord(
            path=rel_path,
            content_hash=h,
            status="ingested",
            summary=result.summary,
            quality=result.quality,
            language=result.language,
            ingested_at=datetime.now(),
        )
    )

    # Normalize concept names against existing canonical names, store linkages
    max_concepts = config.pipeline.max_concepts_per_source
    normalized = _normalize_concepts(result.concepts[:max_concepts], db)
    canonical_names = [name for name, _ in normalized]
    db.delete_concepts_for_source(rel_path)
    db.upsert_concepts(rel_path, canonical_names)
    for canonical, aliases in normalized:
        if aliases:
            db.upsert_aliases(canonical, aliases)

    # Create source summary page in wiki/sources/ (no extra LLM call)
    try:
        _create_source_summary_page(path, meta, result, config, body=body)
    except Exception as e:
        log.warning("Source summary page failed for %s: %s", path.name, e)

    log.info(
        "Ingested: %s (quality=%s, concepts=%s)",
        path.name,
        result.quality,
        [c.name for c in result.concepts[:3]],
    )
    emit_event(
        config,
        event_type="function_timing",
        function_name="ingest_note",
        stage="ingest",
        model=config.models.fast,
        success=True,
        outcome="ingested",
        duration_ms=round((time.monotonic() - fn_t0) * 1000.0, 2),
        note=path.name,
    )
    return result


def ingest_all(
    config: Config,
    client: LLMClientProtocol,
    db: StateDB,
    rag=None,
    force: bool = False,
) -> list[tuple[Path, AnalysisResult | None]]:
    """Ingest all markdown files in raw/, including grouped PDF conversions."""
    raw_files = collect_ingest_paths(config)
    # Snapshot concept names once before loop (for consistent prompt context)
    existing_topics = db.list_all_concept_names()
    results = []
    for path in sorted(raw_files):
        result = ingest_note(
            path=path,
            config=config,
            client=client,
            db=db,
            rag=rag,
            existing_topics=existing_topics,
            force=force,
        )
        results.append((path, result))

    for doc_dir in _find_document_source_dirs(config):
        try:
            _write_document_aggregate(doc_dir, config, db)
        except Exception as e:
            log.warning("Document aggregate failed for %s: %s", doc_dir.name, e)

    return results
